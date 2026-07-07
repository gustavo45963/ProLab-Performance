from pptx import Presentation
from pptx.util import Inches, Pt

def create_presentation():
    # Create a presentation object
    prs = Presentation()

    # 1. Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "ProLab Performance"
    subtitle.text = "Aplicação Web de Treino e Performance"

    # 2. Visão Geral
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Visão Geral"
    tf = body_shape.text_frame
    tf.text = "Recursos principais da plataforma:"

    bullet_points = [
        "Planos de treino por objetivo e nível",
        "Calculadora de IMC integrada",
        "Temporizador HIIT para treinos intervalados",
        "Checklist semanal de acompanhamento",
        "Histórico de evolução e medidas",
        "Contas de utilizador com persistência no servidor"
    ]
    for point in bullet_points:
        p = tf.add_paragraph()
        p.text = point
        p.level = 1

    # 3. Stack Tecnológico
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Stack Tecnológico"
    tf = body_shape.text_frame
    tf.text = "Tecnologias utilizadas no projeto:"

    bullet_points = [
        "Frontend: HTML, CSS (design system próprio), JavaScript (ES Modules, sem frameworks)",
        "Backend: PHP 8.3 (Arquitetura em camadas: Router -> Controllers -> Repositories)",
        "Base de Dados: PostgreSQL (Produção/Render) e SQLite/MySQL (Desenvolvimento)",
        "Deploy: Docker (imagem php:8.3-apache) hospedado no Render"
    ]
    for point in bullet_points:
        p = tf.add_paragraph()
        p.text = point
        p.level = 1

    # 4. Arquitetura do Sistema
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Arquitetura do Sistema"
    tf = body_shape.text_frame
    tf.text = "Estrutura do projeto:"

    bullet_points = [
        "public/: Docroot único (exposta ao browser), contém a SPA (index.html) e front controller (api.php)",
        "src/: Código da aplicação (Bootstrap, Core, Database, Repositories, Controllers, Content)",
        "Database Schema: Gerado automaticamente no primeiro arranque (migração idempotente)"
    ]
    for point in bullet_points:
        p = tf.add_paragraph()
        p.text = point
        p.level = 1

    # 5. Endpoints da API
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Endpoints da API REST"
    tf = body_shape.text_frame
    tf.text = "A API gere toda a lógica e dados via JSON:"

    bullet_points = [
        "/api/auth/* : Autenticação (registo, login, logout, sessão)",
        "/api/profile : Gestão de perfil de utilizador",
        "/api/measurements : Histórico de medidas do utilizador",
        "/api/plans : Catálogo de planos de treino e rotinas",
        "/api/checklist : Acompanhamento e checklist semanal"
    ]
    for point in bullet_points:
        p = tf.add_paragraph()
        p.text = point
        p.level = 1

    # 6. Segurança
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Segurança"
    tf = body_shape.text_frame
    tf.text = "Boas práticas implementadas:"

    bullet_points = [
        "Passwords com hash forte (bcrypt/argon2) via password_hash",
        "Sessões protegidas (HttpOnly, SameSite=Lax, Secure em HTTPS)",
        "Prevenção contra SQL Injection usando Prepared Statements (PDO)",
        "Validação de entradas nos endpoints (código HTTP 422 para erros)",
        "Docroot restrito a 'public/', mantendo código e dados fora do browser"
    ]
    for point in bullet_points:
        p = tf.add_paragraph()
        p.text = point
        p.level = 1

    # Save presentation
    prs.save('ProLab_Performance.pptx')
    print("Presentation generated successfully: ProLab_Performance.pptx")

if __name__ == '__main__':
    create_presentation()
